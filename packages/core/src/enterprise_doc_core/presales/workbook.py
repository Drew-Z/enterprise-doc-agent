"""Bounded Excel intake. Read with openpyxl; preserve OOXML parts on delivery."""
# ruff: noqa: RUF001

from __future__ import annotations

import base64
import binascii
import hashlib
import posixpath
import re
import warnings
from io import BytesIO
from typing import Any, Literal
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple, range_boundaries
from openpyxl.worksheet.cell_range import CellRange

from enterprise_doc_core.presales.errors import PresalesError
from enterprise_doc_core.presales.export import STATUS_LABELS
from enterprise_doc_core.presales.schemas import PacketView, RowView
from enterprise_doc_core.presales.workbook_schemas import (
    MAX_WORKBOOK_BYTES,
    MAX_WORKBOOK_ROWS,
    WorkbookMapping,
    WorkbookMetadata,
    WorkbookPreview,
    WorkbookQuestion,
    WorkbookSheet,
    WorkbookUpload,
)

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def decode_workbook(payload: WorkbookUpload) -> bytes:
    try:
        content = base64.b64decode(payload.content_base64, validate=True)
    except (ValueError, binascii.Error):
        raise PresalesError("presales_workbook_invalid") from None
    if len(content) > MAX_WORKBOOK_BYTES:
        raise PresalesError("presales_workbook_size")
    return content


def _xml(content: bytes) -> Any:
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    root = etree.fromstring(content, parser=parser)
    if root.getroottree().docinfo.doctype:
        raise PresalesError("presales_workbook_unsupported")
    return root


def _parts(content: bytes, filename: str) -> dict[str, bytes]:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise PresalesError("presales_workbook_size")
    if not filename.lower().endswith(".xlsx") or any(c in filename for c in "/\\\x00\r\n"):
        raise PresalesError("presales_workbook_unsupported")
    with ZipFile(BytesIO(content)) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        if len(infos) > 256 or sum(info.file_size for info in infos) > 20 * 1024 * 1024:
            raise PresalesError("presales_workbook_size")
        if len(set(names)) != len(names) or any(
            info.flag_bits & 1
            or info.compress_type not in {ZIP_STORED, ZIP_DEFLATED}
            or info.filename.startswith("/")
            or "\\" in info.filename
            or ".." in info.filename.split("/")
            for info in infos
        ):
            raise PresalesError("presales_workbook_invalid")
        parts = {info.filename: archive.read(info) for info in infos}
    forbidden = (
        "vbaproject",
        "externallinks/",
        "embeddings/",
        "activex/",
        "_xmlsignatures/",
        "macrosheets/",
    )
    if any(any(word in name.lower() for word in forbidden) for name in parts):
        raise PresalesError("presales_workbook_unsupported")
    cell_count = 0
    for name, data in parts.items():
        if not name.endswith((".xml", ".rels")):
            continue
        root = _xml(data)
        if sum(1 for _ in root.iter()) > 200_000:
            raise PresalesError("presales_workbook_size")
        if name.endswith(".rels") and any(item.get("TargetMode") == "External" for item in root):
            raise PresalesError("presales_workbook_unsupported")
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            cell_count += len(root.findall(f".//{{{NS}}}c"))
            coordinates: set[str] = set()
            row_numbers: set[str] = set()
            for sheet_row in root.findall(f"{{{NS}}}sheetData/{{{NS}}}row"):
                number = sheet_row.get("r", "")
                if number in row_numbers or not number.isdecimal() or not 1 <= int(number) <= 10000:
                    raise PresalesError("presales_workbook_invalid")
                row_numbers.add(number)
                for cell in sheet_row.findall(f"{{{NS}}}c"):
                    coordinate = cell.get("r", "")
                    if coordinate in coordinates or coordinate_to_tuple(coordinate)[0] != int(
                        number
                    ):
                        raise PresalesError("presales_workbook_invalid")
                    coordinates.add(coordinate)
            # Bound physical coordinates before openpyxl materializes cells.
            for cell in root.findall(f".//{{{NS}}}c"):
                row, column = coordinate_to_tuple(cell.get("r", ""))
                if row > 10_000 or column > 256:
                    raise PresalesError("presales_workbook_size")
            merged_area = 0
            for merged in root.findall(f".//{{{NS}}}mergeCell"):
                left, top, right, bottom = range_boundaries(merged.get("ref", ""))
                merged_area += (right - left + 1) * (bottom - top + 1)
                if right > 256 or bottom > 10_000 or merged_area > 50_000:
                    raise PresalesError("presales_workbook_size")
    if cell_count > 50_000:
        raise PresalesError("presales_workbook_size")
    types = _xml(parts["[Content_Types].xml"])
    if not any(
        item.get("PartName") == "/xl/workbook.xml"
        and item.get("ContentType") == XLSX_MIME + ".main+xml"
        for item in types
    ):
        raise PresalesError("presales_workbook_unsupported")
    book = _xml(parts["xl/workbook.xml"])
    protection = book.find(f"{{{NS}}}workbookProtection")
    if book.tag != f"{{{NS}}}workbook" or (protection is not None and bool(protection.attrib)):
        raise PresalesError("presales_workbook_unsupported")
    sheets = book.find(f"{{{NS}}}sheets")
    if sheets is None or not 1 <= len(sheets) <= 20:
        raise PresalesError("presales_workbook_size")
    return parts


def _sheet_path(parts: dict[str, bytes], sheet_name: str) -> str:
    book = _xml(parts["xl/workbook.xml"])
    sheet = next(
        s for s in book.findall(f"{{{NS}}}sheets/{{{NS}}}sheet") if s.get("name") == sheet_name
    )
    rel_id = sheet.get(f"{{{OFFICE_REL}}}id")
    relations = _xml(parts["xl/_rels/workbook.xml.rels"])
    target = next(r.get("Target") for r in relations if r.get("Id") == rel_id)
    path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
    if not path.startswith("xl/worksheets/") or path not in parts:
        raise PresalesError("presales_workbook_unsupported")
    return str(path)


def _questions(
    book: Any, mapping: WorkbookMapping, parts: dict[str, bytes]
) -> list[WorkbookQuestion]:
    if (
        mapping.sheet not in book.sheetnames
        or max(
            column_index_from_string(mapping.question_column),
            column_index_from_string(mapping.answer_column),
        )
        > 256
    ):
        raise PresalesError("presales_workbook_mapping")
    sheet = book[mapping.sheet]
    if sheet.sheet_state != "visible" or sheet.protection.sheet:
        raise PresalesError("presales_workbook_protected")
    xml = _xml(parts[_sheet_path(parts, mapping.sheet)])
    formula_ranges = [
        CellRange(item.get("ref")) for item in xml.findall(f".//{{{NS}}}f") if item.get("ref")
    ]
    raw_cells = {
        cell.get("r"): cell for cell in xml.findall(f"{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c")
    }
    result = []
    for row in range(mapping.first_row, mapping.last_row + 1):
        source = sheet[f"{mapping.question_column}{row}"]
        if source.value is None or not str(source.value).strip():
            continue
        target = sheet[f"{mapping.answer_column}{row}"]
        if (
            source.data_type in {"f", "e"}
            or not isinstance(source.value, str)
            or len(source.value.strip()) > 2000
            or sheet.row_dimensions[row].hidden
            or any(source.coordinate in merged for merged in sheet.merged_cells.ranges)
            or any(
                dimension.hidden and dimension.min <= column_index_from_string(col) <= dimension.max
                for dimension in sheet.column_dimensions.values()
                for col in (mapping.question_column, mapping.answer_column)
            )
            or any(source.coordinate in region for region in formula_ranges)
        ):
            raise PresalesError("presales_workbook_question")
        if (
            target.value is not None
            or target.data_type == "f"
            or any(target.coordinate in merged for merged in sheet.merged_cells.ranges)
            or any(
                target.coordinate in rule.sqref for rule in sheet.data_validations.dataValidation
            )
            or any(target.coordinate in CellRange(table.ref) for table in sheet.tables.values())
            or any(target.coordinate in region for region in formula_ranges)
            or (
                target.coordinate in raw_cells
                and any(key not in {"r", "s", "t"} for key in raw_cells[target.coordinate].attrib)
            )
        ):
            raise PresalesError("presales_workbook_target")
        result.append(
            WorkbookQuestion(
                row=row,
                question_cell=source.coordinate,
                answer_cell=target.coordinate,
                text=source.value.strip(),
            )
        )
        if len(result) > MAX_WORKBOOK_ROWS:
            raise PresalesError("presales_workbook_rows")
    if not result:
        raise PresalesError("presales_workbook_mapping")
    return result


def inspect_workbook(
    content: bytes, filename: str, mapping: WorkbookMapping | None = None
) -> WorkbookPreview:
    try:
        parts = _parts(content, filename)
        with warnings.catch_warnings():
            # Unsupported Excel extensions must be visible, never silently discarded.
            warnings.simplefilter("error")
            book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        try:
            sheets = [
                WorkbookSheet(
                    name=sheet.title,
                    max_row=sheet.max_row,
                    max_column=sheet.max_column,
                    selectable=sheet.sheet_state == "visible" and not sheet.protection.sheet,
                    sample=[
                        {
                            cell.column_letter: str(cell.value)[:160]
                            for cell in row
                            if cell.value is not None
                        }
                        for row in sheet.iter_rows(
                            min_row=1,
                            max_row=min(sheet.max_row, 8),
                            max_col=min(sheet.max_column, 12),
                        )
                    ],
                )
                for sheet in book.worksheets
            ]
            questions = _questions(book, mapping, parts) if mapping else []
            return WorkbookPreview(
                filename=filename,
                sha256=hashlib.sha256(content).hexdigest(),
                sheets=sheets,
                questions=questions,
            )
        finally:
            book.close()
    except PresalesError:
        raise
    except Exception:
        # User-controlled ZIP/XML/library diagnostics may contain document text.
        raise PresalesError("presales_workbook_invalid") from None


def _answer(row: RowView) -> str:
    effective = row.review or row.draft
    if effective is None:
        return "未复核草稿：" + (
            "生成失败，请在工作台重试。" if row.state == "failed" else "尚未完成生成。"
        )
    parts = [
        "已复核" if row.review else "未复核草稿",
        STATUS_LABELS[effective.status],
        effective.answer,
    ]
    if effective.conditions:
        parts.append("响应条件：\n" + "\n".join(effective.conditions))
    if effective.missing_information:
        parts.append("待补材料：\n" + "\n".join(effective.missing_information))
    text = "\n".join(parts)
    if len(text.encode("utf-16-le")) // 2 > 32767:
        raise PresalesError("presales_workbook_answer_length")
    # Preserve literal OOXML escape-looking user/model text, as Excel would decode it.
    return re.sub(r"_x([0-9a-fA-F]{4})_", r"_x005F_x\1_", text)


def export_workbook(
    content: bytes,
    metadata: WorkbookMetadata,
    packet: PacketView,
    mode: Literal["draft", "reviewed"],
) -> bytes:
    if mode == "reviewed" and any(row.review is None or row.draft is None for row in packet.rows):
        raise PresalesError("presales_review_required")
    preview = inspect_workbook(content, metadata.filename, metadata.mapping)
    if (
        preview.sha256 != metadata.sha256
        or metadata.rows != [q.row for q in preview.questions]
        or len(packet.rows) != len(metadata.rows)
    ):
        raise PresalesError("presales_workbook_mismatch")
    try:
        parts = _parts(content, metadata.filename)
        path = _sheet_path(parts, metadata.mapping.sheet)
        root = _xml(parts[path])
        sheet_data = root.find(f"{{{NS}}}sheetData")
        if sheet_data is None:
            raise PresalesError("presales_workbook_invalid")
        for question, response in zip(preview.questions, packet.rows, strict=True):
            if response.requirement.text != question.text:
                raise PresalesError("presales_workbook_mismatch")
            row = next(item for item in sheet_data if item.get("r") == str(question.row))
            cell = next((item for item in row if item.get("r") == question.answer_cell), None)
            if cell is None:
                cell = etree.Element(f"{{{NS}}}c", r=question.answer_cell)
                column = column_index_from_string(metadata.mapping.answer_column)
                following = next(
                    (
                        i
                        for i, item in enumerate(row)
                        if coordinate_to_tuple(item.get("r"))[1] > column
                    ),
                    len(row),
                )
                row.insert(following, cell)
            if any(key not in {"r", "s", "t"} for key in cell.attrib):
                raise PresalesError("presales_workbook_target")
            for item in list(cell):
                if item.tag in {f"{{{NS}}}v", f"{{{NS}}}is"}:
                    cell.remove(item)
            cell.set("t", "inlineStr")
            inline = etree.Element(f"{{{NS}}}is")
            text_node = etree.SubElement(inline, f"{{{NS}}}t")
            text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text_node.text = _answer(response)
            cell.insert(0, inline)
        # Expand the used range only if a newly selected answer column lies outside it.
        dimension = root.find(f"{{{NS}}}dimension")
        if dimension is not None:
            from openpyxl.utils.cell import get_column_letter

            left, top, right, bottom = range_boundaries(dimension.get("ref"))
            answer_col = column_index_from_string(metadata.mapping.answer_column)
            dimension.set(
                "ref",
                f"{get_column_letter(min(left, answer_col))}{min(top, min(metadata.rows))}:"
                f"{get_column_letter(max(right, answer_col))}{max(bottom, max(metadata.rows))}",
            )
        edited = etree.tostring(root, encoding="utf-8")
        output = BytesIO()
        with ZipFile(BytesIO(content)) as original, ZipFile(output, "w") as target:
            target.comment = original.comment
            for info in original.infolist():
                target.writestr(info, edited if info.filename == path else parts[info.filename])
        return output.getvalue()
    except PresalesError:
        raise
    except Exception:
        raise PresalesError("presales_workbook_invalid") from None
