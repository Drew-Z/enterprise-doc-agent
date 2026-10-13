from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from enterprise_doc_core.analysis.service import AnalysisError

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEETS = ("Comparison", "Pending Items", "Sources & Basis")


def _safe(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(item) for item in value)
    elif isinstance(value, dict):
        value = ", ".join(f"{key}: {item}" for key, item in value.items())
    else:
        value = str(value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + value
    return value


def _write_row(sheet: Any, values: list[object]) -> None:
    sheet.append([_safe(value) for value in values])


def _style_sheet(sheet: Any, *, widths: tuple[int, ...]) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def render_analysis_workbook(manifest: dict[str, Any]) -> bytes:
    workbook = Workbook()
    comparison = workbook.active
    comparison.title = SHEETS[0]
    pending = workbook.create_sheet(SHEETS[1])
    basis = workbook.create_sheet(SHEETS[2])

    objects = list(manifest["objects"])
    fields = list(manifest["fields"])
    cells = {(cell["object_id"], cell["field_id"]): cell for cell in manifest["cells"]}
    _write_row(comparison, ["Field", *[item["name"] for item in objects]])
    for field in fields:
        values: list[object] = [field["name"]]
        for obj in objects:
            cell = cells[(obj["id"], field["id"])]
            result = cell.get("result")
            if result is None:
                values.append(
                    f"Unavailable ({cell['status']}: {cell.get('error_code') or 'no result'})"
                )
            else:
                values.append(
                    result.get("value") if result.get("value") is not None else result["reason"]
                )
        _write_row(comparison, values)
    _style_sheet(comparison, widths=(28, *([34] * len(objects))))

    _write_row(
        pending,
        ["Object", "Field", "State", "Reason", "Revision", "Review"],
    )
    for cell in manifest["cells"]:
        result = cell.get("result")
        review = cell.get("review")
        if (
            result is not None
            and result.get("action") == "final"
            and result.get("evidence_state") == "supported"
        ):
            continue
        _write_row(
            pending,
            [
                cell["object_name"],
                cell["field_name"],
                result.get("evidence_state") if result is not None else cell["status"],
                result.get("reason")
                if result is not None
                else cell.get("error_code") or "No result",
                cell.get("revision") or "",
                review.get("decision") if review is not None else "unreviewed",
            ],
        )
    _style_sheet(pending, widths=(24, 24, 18, 60, 12, 20))

    _write_row(
        basis,
        ["Kind", "Object / File", "Field / Location", "Basis", "Identity"],
    )
    for source in manifest["sources"]:
        _write_row(
            basis,
            [
                "Source",
                source["filename"],
                source.get("object_id") or "Common requirement",
                f"Coverage: {source['coverage']}; blocks: {source['block_count']}",
                f"version={source['document_version_id']}; "
                f"generation={source['generation_id']}; "
                f"sha256={source['source_sha256']}",
            ],
        )
    for cell in manifest["cells"]:
        result = cell.get("result")
        if result is None:
            continue
        citations = result.get("evidence", [])
        calculations = result.get("calculations", [])
        if not citations and not calculations:
            _write_row(
                basis,
                [
                    "User provided",
                    cell["object_name"],
                    cell["field_name"],
                    result["reason"],
                    f"revision={cell['revision']}; hash={cell['content_hash']}",
                ],
            )
        for citation in citations:
            locator = citation["locator"]
            _write_row(
                basis,
                [
                    "Evidence",
                    citation["filename"],
                    f"{cell['object_name']} / {cell['field_name']} / {locator}",
                    citation["quote"],
                    f"block={citation['block_id']}; hash={citation['quote_hash']}",
                ],
            )
        for calculation in calculations:
            _write_row(
                basis,
                [
                    "Calculation",
                    cell["object_name"],
                    cell["field_name"],
                    calculation["expression"],
                    f"id={calculation['id']}; formula={calculation['formula_version']}",
                ],
            )
    _style_sheet(basis, widths=(18, 30, 38, 70, 70))

    workbook.properties.title = str(manifest["title"])
    workbook.properties.subject = f"Analysis {manifest['mode']} export"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    content = stream.getvalue()
    _verify_workbook(content, expected_cells=len(manifest["cells"]))
    return content


def _verify_workbook(content: bytes, *, expected_cells: int) -> None:
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
        if tuple(workbook.sheetnames) != SHEETS:
            raise AnalysisError("analysis_artifact_invalid")
        if workbook[SHEETS[0]].max_row - 1 <= 0 or expected_cells <= 0:
            raise AnalysisError("analysis_artifact_invalid")
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise AnalysisError("analysis_artifact_invalid")
    except AnalysisError:
        raise
    except Exception as error:
        raise AnalysisError("analysis_artifact_invalid") from error
    finally:
        if "workbook" in locals():
            workbook.close()


__all__ = ["CONTENT_TYPE", "render_analysis_workbook"]
