from __future__ import annotations

import hashlib
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO
from typing import Literal
from zipfile import ZipFile

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from pydantic.alias_generators import to_camel
from pypdf import PdfReader

from enterprise_doc_core.config import UploadSettings
from enterprise_doc_core.documents.envelope import (
    DocumentEnvelopeViolation,
    validate_office_archive_bytes,
)
from enterprise_doc_core.documents.ingestion import DocumentParseViolation, ParsedSection


class StructureModel(BaseModel):
    model_config = ConfigDict(
        extra="forbid", frozen=True, populate_by_name=True, alias_generator=to_camel
    )


class SourceLocator(StructureModel):
    kind: Literal["txt", "pdf", "docx", "xlsx"]
    page: int | None = Field(default=None, ge=1)
    line: int | None = Field(default=None, ge=1)
    part: str | None = None
    paragraph: int | None = Field(default=None, ge=1)
    table: str | None = None
    row: int | None = Field(default=None, ge=1)
    column: int | None = Field(default=None, ge=1)
    row_span: int = Field(default=1, ge=1)
    column_span: int = Field(default=1, ge=1)
    sheet: str | None = None
    cell: str | None = None
    merged_range: str | None = None
    hidden: bool = False


class StructureIssue(StructureModel):
    code: str
    locator: SourceLocator


class DocumentBlockData(StructureModel):
    index: int = Field(ge=0)
    block_id: str
    kind: Literal["paragraph", "heading", "table_cell"]
    text: str
    original_text: str
    start_offset: int = Field(ge=0)
    end_offset: int = Field(ge=0)
    content_sha256: str
    locator: SourceLocator
    heading: str | None = None
    formula: str | None = None
    cached_value: str | None = None
    number_format: str | None = None


class ParsedDocumentStructure(StructureModel):
    schema_version: Literal["document.structure.v1"] = "document.structure.v1"
    parser_version: Literal[2] = 2
    source_sha256: str
    text: str
    blocks: tuple[DocumentBlockData, ...]
    coverage: Literal["complete", "partial", "unknown"]
    issues: tuple[StructureIssue, ...]

    def sections(self) -> tuple[ParsedSection, ...]:
        return tuple(
            ParsedSection(
                block.text,
                block.start_offset,
                block.end_offset,
                block.locator.page,
                block.heading,
            )
            for block in self.blocks
        )


class StructureLimits(StructureModel):
    max_file_bytes: int = Field(default=10 * 1024**2, ge=1, le=10 * 1024**2)
    max_pages: int = Field(default=100, ge=1, le=100)
    max_cells: int = Field(default=20_000, ge=1, le=20_000)
    max_grid_cells: int = Field(default=200_000, ge=1, le=200_000)
    max_text_chars: int = Field(default=1_000_000, ge=1, le=1_000_000)
    max_blocks: int = Field(default=25_000, ge=1, le=25_000)
    max_table_depth: int = Field(default=8, ge=1, le=8)


@dataclass
class _Builder:
    source_sha256: str
    limits: StructureLimits
    blocks: list[DocumentBlockData] = field(default_factory=list)
    issues: list[StructureIssue] = field(default_factory=list)
    issue_keys: set[StructureIssue] = field(default_factory=set)
    cursor: int = 0
    cells: int = 0
    uncertain_layout: bool = False

    def issue(self, code: str, locator: SourceLocator) -> None:
        value = StructureIssue(code=code, locator=locator)
        if value not in self.issue_keys:
            self.issue_keys.add(value)
            self.issues.append(value)

    def add(
        self,
        text: str,
        locator: SourceLocator,
        *,
        kind: Literal["paragraph", "heading", "table_cell"] = "paragraph",
        heading: str | None = None,
        formula: str | None = None,
        cached_value: str | None = None,
        number_format: str | None = None,
    ) -> None:
        normalized = unicodedata.normalize(
            "NFC", text.replace("\r\n", "\n").replace("\r", "\n")
        ).strip()
        if not normalized:
            return
        if "\x00" in normalized:
            raise DocumentParseViolation("structure_nul", "Document text contains a NUL byte.")
        if len(self.blocks) >= self.limits.max_blocks:
            raise DocumentParseViolation("structure_block_limit", "Document has too many blocks.")
        if self.cursor + len(normalized) > self.limits.max_text_chars:
            raise DocumentParseViolation("structure_text_limit", "Document text exceeds the limit.")
        if kind == "table_cell":
            self.cells += 1
            if self.cells > self.limits.max_cells:
                raise DocumentParseViolation("structure_cell_limit", "Document has too many cells.")
        digest = hashlib.sha256(normalized.encode()).hexdigest()
        identity = f"{self.source_sha256}:2:{len(self.blocks)}:{digest}"
        self.blocks.append(
            DocumentBlockData(
                index=len(self.blocks),
                block_id=hashlib.sha256(identity.encode()).hexdigest(),
                kind=kind,
                text=normalized,
                original_text=text,
                start_offset=self.cursor,
                end_offset=self.cursor + len(normalized),
                content_sha256=digest,
                locator=locator,
                heading=heading,
                formula=formula,
                cached_value=cached_value,
                number_format=number_format,
            )
        )
        self.cursor += len(normalized) + 1

    def finish(self) -> ParsedDocumentStructure:
        return ParsedDocumentStructure(
            source_sha256=self.source_sha256,
            text="\n".join(block.text for block in self.blocks),
            blocks=tuple(self.blocks),
            coverage="partial"
            if self.issues
            else ("unknown" if self.uncertain_layout else "complete"),
            issues=tuple(self.issues),
        )


def parse_document_structure(
    content: bytes, *, extension: str, limits: StructureLimits | None = None
) -> ParsedDocumentStructure:
    """Parse a bounded source into immutable blocks; never evaluate Office content."""
    resolved = limits or StructureLimits()
    if not content or len(content) > resolved.max_file_bytes:
        raise DocumentParseViolation("structure_file_limit", "Document size is outside the limit.")
    extension = "." + extension.lower().strip().lstrip(".")
    builder = _Builder(hashlib.sha256(content).hexdigest(), resolved)
    try:
        if extension in {".docx", ".xlsx"}:
            validate_office_archive_bytes(
                content,
                extension=extension,
                settings=UploadSettings(
                    max_docx_entries=2000,
                    max_docx_declared_uncompressed_bytes=64 * 1024**2,
                    max_docx_member_uncompressed_bytes=16 * 1024**2,
                ),
            )
            with ZipFile(BytesIO(content)) as archive:
                names = archive.namelist()
                if any("vbaproject" in name.lower() for name in names):
                    raise DocumentParseViolation(
                        "structure_macro_unsupported", "Macros are unsupported."
                    )
        if extension == ".txt":
            _txt(content, builder)
        elif extension == ".pdf":
            _pdf(content, builder)
        elif extension == ".docx":
            _docx(content, builder)
        elif extension == ".xlsx":
            _xlsx(content, builder)
        else:
            raise DocumentParseViolation(
                "document_type_unsupported", "Document type is unsupported."
            )
    except DocumentEnvelopeViolation as error:
        raise DocumentParseViolation(error.code, error.message) from error
    except DocumentParseViolation:
        raise
    except Exception as error:
        raise DocumentParseViolation(
            "structure_parse_failed", "Document structure could not be parsed."
        ) from error
    return builder.finish()


def _txt(content: bytes, builder: _Builder) -> None:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentParseViolation(
            "text_decode_failed", "Text document is not valid UTF-8."
        ) from error
    heading = None
    for index, line in enumerate(text.splitlines(), 1):
        is_heading = line.startswith("#") and " " in line
        if is_heading:
            heading = line.lstrip("#").strip()
        builder.add(
            line,
            SourceLocator(kind="txt", line=index),
            kind="heading" if is_heading else "paragraph",
            heading=heading,
        )


def _pdf(content: bytes, builder: _Builder) -> None:
    reader = PdfReader(BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise DocumentParseViolation("structure_encrypted", "Encrypted documents are unsupported.")
    if len(reader.pages) > builder.limits.max_pages:
        raise DocumentParseViolation("structure_page_limit", "Document has too many pages.")
    builder.uncertain_layout = True
    for index, page in enumerate(reader.pages, 1):
        locator = SourceLocator(kind="pdf", page=index)
        text = page.extract_text() or ""
        if not text.strip():
            builder.issue("needs_ocr", locator)
        else:
            builder.add(text, locator)
            resources = page.get("/Resources")
            if resources is not None and resources.get_object().get("/XObject"):
                builder.issue("pdf_visual_content_unverified", locator)


def _docx(content: bytes, builder: _Builder) -> None:
    document = DocxDocument(BytesIO(content))
    heading: str | None = None
    for index, item in enumerate(document.iter_inner_content(), 1):
        if isinstance(item, Paragraph):
            is_heading = item.style is not None and item.style.name.lower().startswith("heading")
            if is_heading:
                heading = item.text.strip()
            builder.add(
                item.text,
                SourceLocator(kind="docx", part="word/document.xml", paragraph=index),
                kind="heading" if is_heading else "paragraph",
                heading=heading,
            )
        else:
            _docx_table(item, builder, "word/document.xml", f"body/{index}", heading, 1)
    seen_parts = {"word/document.xml"}
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            if container.is_linked_to_previous:
                continue
            part_name = str(container.part.partname).lstrip("/")
            if part_name in seen_parts:
                continue
            seen_parts.add(part_name)
            for index, item in enumerate(container.iter_inner_content(), 1):
                if isinstance(item, Paragraph):
                    builder.add(
                        item.text, SourceLocator(kind="docx", part=part_name, paragraph=index)
                    )
                else:
                    _docx_table(item, builder, part_name, f"body/{index}", None, 1)
    unsupported = {
        "drawing",
        "pict",
        "object",
        "txbxContent",
        "ins",
        "del",
        "footnoteReference",
        "endnoteReference",
        "altChunk",
        "sdt",
    }
    for part in document.part.package.parts:
        part_name = str(part.partname).lstrip("/")
        element = getattr(part, "element", None)
        if part_name in seen_parts and element is not None:
            if any(str(node.tag).rsplit("}", 1)[-1] in unsupported for node in element.iter()):
                builder.issue(
                    "docx_unsupported_content", SourceLocator(kind="docx", part=part_name)
                )


def _docx_table(
    table: Table, builder: _Builder, part: str, path: str, heading: str | None, depth: int
) -> None:
    if depth > builder.limits.max_table_depth:
        raise DocumentParseViolation("structure_table_depth", "Nested table limit exceeded.")
    seen: set[object] = set()
    table_rows = list(table.rows)
    rows = [list(row.cells) for row in table_rows]
    if sum(len(row) for row in rows) > builder.limits.max_grid_cells:
        raise DocumentParseViolation("structure_grid_limit", "Table grid exceeds the limit.")
    positions_by_cell: dict[object, list[tuple[int, int]]] = {}
    for r, values in enumerate(rows, 1):
        for c, cell in enumerate(values, table_rows[r - 1].grid_cols_before + 1):
            positions_by_cell.setdefault(cell._tc, []).append((r, c))
    for row_index, row in enumerate(rows, 1):
        for column_index, cell in enumerate(row, table_rows[row_index - 1].grid_cols_before + 1):
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            positions = positions_by_cell[cell._tc]
            locator = SourceLocator(
                kind="docx",
                part=part,
                table=path,
                row=row_index,
                column=column_index,
                row_span=max(p[0] for p in positions) - row_index + 1,
                column_span=max(p[1] for p in positions) - column_index + 1,
            )
            for index, item in enumerate(cell.iter_inner_content(), 1):
                if isinstance(item, Paragraph):
                    builder.add(
                        item.text,
                        locator.model_copy(update={"paragraph": index}),
                        kind="table_cell",
                        heading=heading,
                    )
                else:
                    _docx_table(
                        item,
                        builder,
                        part,
                        f"{path}/{row_index}/{column_index}/{index}",
                        heading,
                        depth + 1,
                    )


def _xlsx(content: bytes, builder: _Builder) -> None:
    workbook = load_workbook(BytesIO(content), read_only=False, data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=False)
    try:
        grid_cells = sum(sheet.max_row * sheet.max_column for sheet in workbook.worksheets)
        if grid_cells > builder.limits.max_grid_cells:
            raise DocumentParseViolation(
                "structure_grid_limit", "Worksheet grid exceeds the limit."
            )
        with ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            if any(
                name.startswith(("xl/externalLinks/", "xl/drawings/", "xl/embeddings/"))
                for name in names
            ):
                builder.issue("xlsx_unsupported_content", SourceLocator(kind="xlsx"))
        for sheet in workbook.worksheets:
            hidden_columns = {
                index
                for dimension in sheet.column_dimensions.values()
                if dimension.hidden
                for index in range(dimension.min, dimension.max + 1)
            }
            merged_at_origin = {
                (item.min_row, item.min_col): str(item) for item in sheet.merged_cells.ranges
            }
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    merge = merged_at_origin.get((cell.row, cell.column))
                    locator = SourceLocator(
                        kind="xlsx",
                        sheet=sheet.title,
                        cell=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        merged_range=merge,
                        hidden=(
                            sheet.sheet_state != "visible"
                            or bool(sheet.row_dimensions[cell.row].hidden)
                            or cell.column in hidden_columns
                        ),
                    )
                    formula = str(cell.value) if cell.data_type == "f" else None
                    cached_value = cached[sheet.title][cell.coordinate].value if formula else None
                    if formula:
                        builder.issue(
                            "formula_cache_missing"
                            if cached_value is None
                            else "formula_cache_unverified",
                            locator,
                        )
                    if cell.data_type == "e":
                        builder.issue("spreadsheet_error_value", locator)
                    builder.add(
                        str(cell.value),
                        locator,
                        kind="table_cell",
                        heading=sheet.title,
                        formula=formula,
                        cached_value=str(cached_value) if cached_value is not None else None,
                        number_format=cell.number_format,
                    )
    finally:
        workbook.close()
        cached.close()
