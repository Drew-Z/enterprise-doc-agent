import base64
import json
from io import BytesIO
from uuid import UUID

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from enterprise_doc_core.billing.models import UsageEvent
from enterprise_doc_core.presales.background import BackgroundGeneration
from enterprise_doc_core.presales.gateway import OpenAICompatiblePresalesGateway
from enterprise_doc_core.presales.models import PresalesProviderCall
from enterprise_doc_core.presales.schemas import SourceInput
from enterprise_doc_core.presales.workbook import inspect_workbook
from tests.browser_sessions.conftest import browser_db as browser_db
from tests.presales.fixtures import add_chunk, add_document
from tests.presales.test_presales_background_integration import background as background
from tests.presales.test_presales_background_integration import valid_response
from tests.presales.test_workbook import mapping

pytestmark = pytest.mark.integration


@pytest.mark.parametrize("primary_fault", [None, "502", "duplicate_question", "many_gaps"])
async def test_question_generation_review_reload_and_workbook_keep_draft_and_accounting(
    background, primary_fault
):
    b = background
    needs_fallback = primary_fault in {"502", "duplicate_question"}
    version, generation = await add_document(b.sessions, b.context)
    await add_chunk(
        b.sessions,
        b.context,
        version,
        generation,
        "Retention: 启用保留功能前须配置策略。策略尚未配置。",
    )
    b.payload = b.payload.model_copy(
        update={
            "sources": [
                *b.payload.sources,
                SourceInput(version_id=version, applicability="Current purchase"),
            ]
        }
    )
    b.settings.primary_question_assessment = True
    b.settings.automatic_failover_enabled = needs_fallback
    calls = []

    def respond(request):
        body = json.loads(request.content)
        calls.append((request.url.host, body))
        if primary_fault == "502" and request.url.host == "primary.invalid":
            return httpx.Response(502)
        if request.url.host == "fallback.invalid":
            return valid_response(request)
        wire = json.loads(body["messages"][1]["content"])
        refs = {s["text"]: {"spanId": s["spanId"]} for s in wire["spans"]}
        cite = next(e["citationId"] for e in wire["evidence"] if "策略尚未配置" in e["text"])
        value = {
            "rules": [
                {
                    "proposition": "完成保留策略配置",
                    "requiredBy": [refs["Retention: 启用保留功能前须配置策略。"]],
                }
            ],
            "assessments": [
                {
                    "ruleIndex": 0,
                    "state": "unmet",
                    "evidence": [refs["策略尚未配置。"]],
                    "summary": "本订单策略尚未配置。",
                    "nextAction": "完成所需策略配置。",
                }
            ],
            "responses": [
                {
                    "requirementPartId": p["requirementPartId"],
                    "answer": "策略尚未配置\uff0c当前不能启用。",
                    "citations": [{"citationId": cite}],
                    "missingInformation": [],
                }
                for p in wire["requirementParts"]
            ],
            "status": "conditional",
            "conclusion": "完成策略配置后方可启用。",
        }
        if primary_fault == "duplicate_question":
            value["responses"].append(dict(value["responses"][0]))
        elif primary_fault == "many_gaps":
            for index, response in enumerate(value["responses"]):
                response["missingInformation"] = [
                    f"请确认第{index}部分的第{gap}项后续安排。" for gap in range(7)
                ]
        return httpx.Response(
            200,
            json={
                "model": "gpt-6-luna",
                "id": "controlled-question",
                "choices": [{"finish_reason": "stop", "message": {"content": json.dumps(value)}}],
                "usage": {"total_tokens": 50},
            },
        )

    root = b.service.generation.gateway.source_settings
    if needs_fallback:
        root = root.model_copy(
            update={
                "fallback_provider": root.provider,
                "fallback_base_url": "https://fallback.invalid/v1",
                "fallback_api_key": root.api_key,
                "fallback_model_name": "fallback-test-model",
            }
        )
    transport = httpx.MockTransport(respond)
    primary = OpenAICompatiblePresalesGateway(
        root, presales_settings=b.settings, transport=transport
    )
    b.service.generation.gateway = primary
    gateways = {"primary": primary}
    if needs_fallback:
        gateways["fallback"] = OpenAICompatiblePresalesGateway(
            root,
            presales_settings=b.settings.model_copy(update={"model_route": "fallback"}),
            transport=transport,
        )
    book = Workbook()
    sheet = book.active
    sheet.title = "技术要求"
    sheet.append(["编号", "问题", "答复", "保留内容"])
    question = "Retention. Retention." if primary_fault == "many_gaps" else "Retention"
    sheet.append(["R1", question, None, "不可改写"])
    buffer = BytesIO()
    book.save(buffer)
    content = buffer.getvalue()
    headers = {"Authorization": "Bearer owner", "Idempotency-Key": "question-import"}
    created = await b.api.post(
        "/api/presales/workbooks",
        headers=headers,
        json={
            "filename": "questionnaire.xlsx",
            "contentBase64": base64.b64encode(content).decode(),
            "title": "Question protocol delivery",
            "sources": [s.model_dump(mode="json", by_alias=True) for s in b.payload.sources],
            "mapping": mapping(1).model_dump(mode="json", by_alias=True),
            "confirmedSha256": inspect_workbook(content, "questionnaire.xlsx").sha256,
        },
    )
    assert created.status_code == 201, created.text
    packet = created.json()
    url = f"/api/presales/{packet['id']}"
    row_url = url + f"/rows/{packet['rows'][0]['id']}"
    admitted = await b.api.post(
        row_url + "/generate",
        headers={**headers, "Idempotency-Key": "question-generate"},
        json={"executionMode": "auto"},
    )
    assert admitted.status_code == 202, admitted.text
    assert calls == []
    worker = BackgroundGeneration(b.service.generation, gateways)
    assert await worker.run_once("question-worker")
    assert not await worker.run_once("question-no-replay")
    loaded = await b.api.get(url, headers=headers)
    row = loaded.json()["rows"][0]
    assert row["state"] == "drafted", json.dumps(row, ensure_ascii=False)
    assert row["attempts"][0]["executionPolicy"]["routes"][0]["promptVersion"] == "presales.v24"
    if not needs_fallback:
        assert row["draft"]["prerequisites"][0]["state"] == "unmet"
        assert "完成所需策略配置。" in row["draft"]["answer"]
    if primary_fault == "many_gaps":
        expected = [
            f"请确认第{index}部分的第{gap}项后续安排。" for index in range(2) for gap in range(7)
        ]
        assert row["draft"]["missingInformation"] == expected
    draft = row["draft"]
    response_text = "人工核对后的交付说明\uff1a" + draft["answer"]
    review = {
        "expectedRevision": row["revision"],
        "status": draft["status"],
        "answer": response_text,
        "conditions": draft["conditions"],
        "prerequisites": draft["prerequisites"],
        "prerequisiteChanges": {
            "origins": list(range(len(draft["prerequisites"] or []))),
            "excludedIndexes": [],
        },
        "missingInformation": draft["missingInformation"],
        "note": "测试复核保留原始模型草稿。",
    }
    reviewed = await b.api.put(
        row_url + "/review", headers={**headers, "Idempotency-Key": "question-review"}, json=review
    )
    assert reviewed.status_code == 200, reviewed.text
    refreshed = (await b.api.get(url, headers=headers)).json()["rows"][0]
    assert refreshed["draft"] == draft and len(refreshed["reviewHistory"]) == 1
    assert refreshed["review"]["missingInformation"] == draft["missingInformation"]
    assert refreshed["review"]["actorId"] == str(b.context.actor_id)
    downloaded = await b.api.get(url + "/workbook?mode=reviewed", headers=headers)
    assert downloaded.status_code == 200
    exported = load_workbook(BytesIO(downloaded.content))
    assert response_text in exported["技术要求"]["C2"].value
    assert exported["技术要求"]["D2"].value == "不可改写"
    csv = await b.api.get(url + "/export?mode=reviewed", headers=headers)
    assert csv.status_code == 200 and response_text in csv.content.decode("utf-8-sig")
    if primary_fault == "many_gaps":
        assert "\n".join(expected) in exported["技术要求"]["C2"].value
        assert all(gap in csv.content.decode("utf-8-sig") for gap in expected)
    assert len(calls) == (2 if needs_fallback else 1)
    async with b.sessions() as session:
        operation = UUID(row["attempts"][0]["id"])
        records = (
            await session.scalars(
                select(PresalesProviderCall)
                .where(PresalesProviderCall.operation_id == operation)
                .order_by(PresalesProviderCall.number)
            )
        ).all()
        assert [c.state for c in records] == (
            ["failed", "succeeded"] if needs_fallback else ["succeeded"]
        )
        events = (
            await session.scalars(select(UsageEvent).where(UsageEvent.operation_id == operation))
        ).all()
        assert len([e for e in events if e.event_type == "consume"]) == 1
