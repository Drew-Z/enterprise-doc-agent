from datetime import UTC, datetime
from io import BytesIO
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from enterprise_doc_core.presales.errors import PresalesError
from enterprise_doc_core.presales.schemas import PacketView, RequirementInput, RowView, SavedDraft
from enterprise_doc_core.presales.workbook import inspect_workbook
from enterprise_doc_core.presales.workbook_schemas import WorkbookMapping, WorkbookMetadata


def questionnaire(count=2):
    book = Workbook()
    sheet = book.active
    sheet.title = "技术要求"
    sheet.append(["编号", "客户问题", "供应商答复", "原有公式"])
    for number in range(1, count + 1):
        sheet.append([f"Q-{number}", f"Retention {number}?", None, f"=A{number + 1}"])
        sheet.cell(number + 1, 3).fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.column_dimensions["B"].width = 45
    sheet.column_dimensions["C"].width = 60
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:D{count + 1}"
    book.create_sheet("采购说明")["A1"] = "请保留本页。"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def mapping(count=2):
    return WorkbookMapping(
        sheet="技术要求", question_column="B", answer_column="C", first_row=2, last_row=count + 1
    )


def test_preview_maps_real_workbook_and_keeps_original_row_locations():
    result = inspect_workbook(questionnaire(13), "客户要求.xlsx", mapping(13))
    assert [sheet.name for sheet in result.sheets] == ["技术要求", "采购说明"]
    assert len(result.questions) == 13
    assert result.questions[0].question_cell == "B2"
    assert result.questions[0].answer_cell == "C2"
    assert result.questions[-1].text == "Retention 13?"
    assert result.questions[-1].row == 14
    assert len(result.sha256) == 64


def test_export_preserves_original_parts_styles_formulas_and_literal_answer():
    from enterprise_doc_core.presales.workbook import export_workbook

    content = questionnaire()
    preview = inspect_workbook(content, "客户要求.xlsx", mapping())
    packet = PacketView(
        id=uuid4(),
        title="Questionnaire",
        created_at=datetime.now(UTC),
        row_count=2,
        sources=[],
        rows=[
            RowView(
                id=uuid4(),
                requirement=RequirementInput(key=f"X{q.row}", text=q.text),
                revision=1 if index == 0 else 0,
                state="drafted" if index == 0 else "failed",
                draft=SavedDraft(
                    status="conditional",
                    answer='=HYPERLINK("https://invalid")',
                    conditions=["购买服务"],
                    missing_information=["确认范围"],
                    citations=[],
                    retrieval=[],
                )
                if index == 0
                else None,
                review=None,
                review_history=[],
                attempts=[],
            )
            for index, q in enumerate(preview.questions)
        ],
    )
    metadata = WorkbookMetadata(
        filename=preview.filename, sha256=preview.sha256, mapping=mapping(), rows=[2, 3]
    )
    output = export_workbook(content, metadata, packet, "draft")
    with ZipFile(BytesIO(content)) as before, ZipFile(BytesIO(output)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert before.read(name) == after.read(name)
    book = load_workbook(BytesIO(output))
    assert book["技术要求"]["C2"].data_type == "s"
    assert "=HYPERLINK" in book["技术要求"]["C2"].value
    assert "未复核草稿" in book["技术要求"]["C2"].value
    assert "购买服务" in book["技术要求"]["C2"].value
    assert "确认范围" in book["技术要求"]["C2"].value
    assert "生成失败" in book["技术要求"]["C3"].value
    assert book["技术要求"]["C2"].fill.fgColor.rgb == "00FFF2CC"
    assert book["技术要求"]["D2"].value == "=A2"
    assert book["技术要求"].freeze_panes == "C2"
    assert book["采购说明"]["A1"].value == "请保留本页。"
    with pytest.raises(PresalesError, match="presales_review_required"):
        export_workbook(content, metadata, packet, "reviewed")


@pytest.mark.parametrize(
    "kind,code",
    [
        ("occupied", "target"),
        ("formula", "target"),
        ("merged", "target"),
        ("protected", "protected"),
        ("question_formula", "question"),
    ],
)
def test_import_rejects_unsafe_target_or_question(kind, code):
    book = load_workbook(BytesIO(questionnaire()))
    sheet = book.active
    if kind == "occupied":
        sheet["C2"] = "Customer answer"
    if kind == "formula":
        sheet["C2"] = '=IF(A1="","",A1)'
    if kind == "merged":
        sheet.merge_cells("C2:C3")
    if kind == "protected":
        sheet.protection.sheet = True
    if kind == "question_formula":
        sheet["B2"] = "=A2"
    output = BytesIO()
    book.save(output)
    with pytest.raises(PresalesError, match=f"presales_workbook_{code}"):
        inspect_workbook(output.getvalue(), "customer.xlsx", mapping())


def test_import_rejects_invalid_zip_extension_and_too_many_questions():
    with pytest.raises(PresalesError, match="presales_workbook_invalid"):
        inspect_workbook(b"customer secret body", "customer.xlsx")
    with pytest.raises(PresalesError, match="presales_workbook_unsupported"):
        inspect_workbook(questionnaire(), "customer.xlsm")
    with pytest.raises(PresalesError, match="presales_workbook_rows"):
        inspect_workbook(questionnaire(121), "customer.xlsx", mapping(121))


def replace_part(content, path, transform):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item)
            target.writestr(item, transform(value) if item.filename == path else value)
    return output.getvalue()


def test_rejects_duplicate_cells_before_mapping_and_grouped_hidden_columns():
    import copy

    from lxml import etree

    def duplicate(xml):
        root = etree.fromstring(xml)
        row = root.find("{*}sheetData/{*}row[@r='2']")
        row.append(copy.deepcopy(row[1]))
        return etree.tostring(root)

    duplicated = replace_part(questionnaire(), "xl/worksheets/sheet1.xml", duplicate)
    with pytest.raises(PresalesError, match="presales_workbook_invalid"):
        inspect_workbook(duplicated, "customer.xlsx", mapping())
    book = load_workbook(BytesIO(questionnaire()))
    book.active.column_dimensions.group("A", "D", hidden=True)
    output = BytesIO()
    book.save(output)
    with pytest.raises(PresalesError, match="presales_workbook_question"):
        inspect_workbook(output.getvalue(), "customer.xlsx", mapping())


def test_rejects_answer_inside_array_formula_range():
    from openpyxl.worksheet.formula import ArrayFormula

    book = load_workbook(BytesIO(questionnaire()))
    book.active["C1"] = ArrayFormula(ref="C1:C3", text="=A1:A3")
    output = BytesIO()
    book.save(output)
    with pytest.raises(PresalesError, match="presales_workbook_target"):
        inspect_workbook(output.getvalue(), "customer.xlsx", mapping())
