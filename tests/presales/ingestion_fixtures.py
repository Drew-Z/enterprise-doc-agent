"""Small synthetic files for the real upload/worker browser acceptance run."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


@dataclass(frozen=True, slots=True)
class UploadFixture:
    name: str
    media_type: str
    content: bytes
    excerpt: str
    heading: str | None = None
    page_number: int | None = None

    def metadata(self) -> dict[str, object]:
        return {
            "name": self.name,
            "mediaType": self.media_type,
            "sizeBytes": len(self.content),
            "sha256": hashlib.sha256(self.content).hexdigest(),
            "excerpt": self.excerpt,
            "heading": self.heading,
            "pageNumber": self.page_number,
        }


def build_upload_fixtures() -> dict[str, UploadFixture]:
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=595, height=842)
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 14 Tf 48 780 Td (Export is available in CSV format.) Tj ET")
    page[NameObject("/Contents")] = stream
    pdf_bytes = BytesIO()
    pdf.write(pdf_bytes)

    docx = DocxDocument()
    docx.add_heading("备份策略", level=1)
    docx.add_paragraph("Backups are retained for 7 days.")
    docx_bytes = BytesIO()
    docx.save(docx_bytes)

    return {
        "txt": UploadFixture(
            "合成上传-保留策略.txt",
            "text/plain",
            "# 数据保留策略\nRetention is 30 days.\n".encode(),
            "Retention is 30 days.",
            heading="数据保留策略",
        ),
        "pdf": UploadFixture(
            "合成上传-导出说明.pdf",
            "application/pdf",
            pdf_bytes.getvalue(),
            "Export is available in CSV format.",
            page_number=1,
        ),
        "docx": UploadFixture(
            "合成上传-备份策略.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx_bytes.getvalue(),
            "Backups are retained for 7 days.",
            heading="备份策略",
        ),
        # A valid PDF signature passes upload envelope validation. Missing PDF
        # structure must then fail in the real parser, without injected faults.
        "broken_pdf": UploadFixture(
            "合成上传-损坏文件.pdf", "application/pdf", b"%PDF-1.7\ninvalid body\n%%EOF\n", ""
        ),
        "repaired_pdf": UploadFixture(
            "合成上传-修复文件.pdf",
            "application/pdf",
            pdf_bytes.getvalue(),
            "Export is available in CSV format.",
            page_number=1,
        ),
    }
