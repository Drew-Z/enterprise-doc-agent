from __future__ import annotations

import hashlib
from contextlib import asynccontextmanager
from datetime import UTC, datetime, timedelta
from io import BytesIO
from typing import Any

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from enterprise_doc_core.analysis.execution_models import AnalysisArtifact
from enterprise_doc_core.analysis.review import AnalysisReviewDeliveryService
from enterprise_doc_core.analysis.review_schemas import (
    ArtifactApprovalInput,
    ExportInput,
    ManualRevisionInput,
    ReviewInput,
)
from enterprise_doc_core.analysis.service import AnalysisError
from enterprise_doc_core.documents.models import Document
from enterprise_doc_core.object_store import ArtifactObject, ObjectHead, PresignedObjectDownload
from tests.analysis.test_analysis_execution_integration import claim_run, response, setup

pytestmark = pytest.mark.integration


class MemoryArtifactStore:
    def __init__(self) -> None:
        self.objects: dict[tuple[str, str], tuple[bytes, str, dict[str, str]]] = {}
        self.presigns = 0

    async def put_object(
        self,
        *,
        bucket: str,
        key: str,
        body: bytes,
        content_type: str,
        metadata: dict[str, str] | None = None,
    ) -> ArtifactObject:
        digest = hashlib.sha256(body).hexdigest()
        stored_metadata = {"sha256": digest, **(metadata or {})}
        self.objects[(bucket, key)] = (body, content_type, stored_metadata)
        return ArtifactObject(
            bucket=bucket,
            key=key,
            size_bytes=len(body),
            content_sha256=digest,
            content_type=content_type,
            etag='"memory"',
        )

    async def head_object(self, *, bucket: str, key: str) -> ObjectHead:
        body, content_type, metadata = self.objects[(bucket, key)]
        return ObjectHead(
            size_bytes=len(body),
            etag='"memory"',
            checksum_sha256_b64=None,
            content_type=content_type,
            metadata=metadata,
        )

    async def read_object(self, *, bucket: str, key: str, expected_size: int) -> bytes:
        body = self.objects[(bucket, key)][0]
        assert len(body) == expected_size
        return body

    async def delete_object(self, *, bucket: str, key: str) -> None:
        self.objects.pop((bucket, key), None)

    async def presign_get(
        self, *, bucket: str, key: str, expires_in_seconds: int
    ) -> PresignedObjectDownload:
        assert (bucket, key) in self.objects
        self.presigns += 1
        return PresignedObjectDownload(
            url=f"https://download.test/{key}", expires_in_seconds=expires_in_seconds
        )

    async def close(self) -> None:
        return None


class RollbackSecondTransaction:
    def __init__(self, sessions: Any) -> None:
        self.sessions = sessions
        self.transactions = 0

    def __call__(self):
        return self.sessions()

    @asynccontextmanager
    async def begin(self):
        self.transactions += 1
        async with self.sessions() as session, session.begin():
            yield session
            if self.transactions == 2:
                raise RuntimeError("simulated finalization rollback")


async def test_manual_revision_and_review_are_immutable_exact_and_replay_safe(browser_db) -> None:
    async with httpx.AsyncClient(transport=httpx.MockTransport(response)) as provider:
        execution, kwargs, _ = await setup(browser_db.sessions, provider)
        run = await execution.create_run(**kwargs)
        runtime, claim = await claim_run(browser_db.sessions, run)
        await execution.execute(claim)
        await runtime.succeed(claim)

        service = AnalysisReviewDeliveryService(
            session_factory=browser_db.sessions,
            gateway=execution.gateway,
        )
        before = await service.matrix(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], run_id=run.id
        )
        unit = before.units[0]
        assert unit.current_revision == 1
        assert unit.result is not None and unit.result.value_origin == "document_evidence"

        request = ManualRevisionInput(
            action="final",
            value='=HYPERLINK("https://example.invalid")',
            evidence_state="insufficient",
            reason="Confirmed directly by the buyer.",
            value_origin="user_provided",
        )
        changed = await service.create_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            expected_revision=1,
            idempotency_key="manual-value-v2",
            request=request,
        )
        replay = await service.create_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            expected_revision=1,
            idempotency_key="manual-value-v2",
            request=request,
        )
        assert changed.id == replay.id
        assert changed.revision == 2
        assert changed.author_type == "user"
        assert changed.result.value_origin == "user_provided"
        assert changed.result.evidence == ()

        reviewed = await service.review_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            idempotency_key="approve-v2",
            request=ReviewInput(
                revision=2,
                content_hash=changed.content_hash,
                decision="approved",
                reason="Checked against the buyer confirmation.",
            ),
        )
        reviewed_replay = await service.review_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            idempotency_key="approve-v2",
            request=ReviewInput(
                revision=2,
                content_hash=changed.content_hash,
                decision="approved",
                reason="Checked against the buyer confirmation.",
            ),
        )
        assert reviewed.id == reviewed_replay.id
        assert reviewed.cell_revision_id == changed.id

        latest = await service.create_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            expected_revision=2,
            idempotency_key="manual-value-v3",
            request=request.model_copy(update={"value": "Confirmed static value"}),
        )
        assert latest.revision == 3
        with pytest.raises(AnalysisError, match="analysis_review_revision_conflict"):
            await service.review_cell_revision(
                tenant_id=kwargs["tenant_id"],
                actor_id=kwargs["actor_id"],
                unit_id=unit.id,
                idempotency_key="stale-review",
                request=ReviewInput(
                    revision=2,
                    content_hash=changed.content_hash,
                    decision="approved",
                    reason="This revision is no longer current.",
                ),
            )

        history = await service.cell_revisions(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], unit_id=unit.id
        )
        assert [item.revision for item in history] == [3, 2, 1]
        assert history[1].review is not None and history[1].review.decision == "approved"
        assert history[2].author_type == "model"
        matrix = await service.matrix(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], run_id=run.id
        )
        assert matrix.units[0].current_revision == 3
        assert matrix.units[0].review is None
        assert matrix.units[0].result is not None
        assert matrix.units[0].result.review_state == "unreviewed"


async def test_failed_cell_can_start_with_a_user_provided_revision(browser_db) -> None:
    def invalid_response(_: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={"choices": [{"finish_reason": "stop", "message": {"content": "wrong"}}]},
        )

    async with httpx.AsyncClient(transport=httpx.MockTransport(invalid_response)) as provider:
        execution, kwargs, _ = await setup(browser_db.sessions, provider)
        run = await execution.create_run(**kwargs)
        runtime, claim = await claim_run(browser_db.sessions, run)
        await execution.execute(claim)
        await runtime.succeed(claim)

        service = AnalysisReviewDeliveryService(
            session_factory=browser_db.sessions,
            gateway=execution.gateway,
        )
        before = await service.matrix(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], run_id=run.id
        )
        unit = before.units[0]
        assert unit.status == "failed"
        assert unit.result is None
        assert unit.current_revision is None

        created = await service.create_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            expected_revision=0,
            idempotency_key="manual-first-revision",
            request=ManualRevisionInput(
                action="final",
                value="Confirmed by the buyer after model failure",
                evidence_state="insufficient",
                reason="The model produced no valid result.",
                value_origin="user_provided",
            ),
        )

        assert created.revision == 1
        assert created.supersedes_id is None
        assert created.author_type == "user"
        after = await service.matrix(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], run_id=run.id
        )
        assert after.units[0].status == "failed"
        assert after.units[0].error_code == "analysis_output_invalid"
        assert after.units[0].result is not None
        assert after.units[0].result.value_origin == "user_provided"


async def test_reviewed_xlsx_freezes_matrix_and_download_rechecks_source_access(browser_db) -> None:
    store = MemoryArtifactStore()
    async with httpx.AsyncClient(transport=httpx.MockTransport(response)) as provider:
        execution, kwargs, document_id = await setup(browser_db.sessions, provider)
        run = await execution.create_run(**kwargs)
        runtime, claim = await claim_run(browser_db.sessions, run)
        await execution.execute(claim)
        await runtime.succeed(claim)
        service = AnalysisReviewDeliveryService(
            session_factory=browser_db.sessions,
            gateway=execution.gateway,
            artifact_store=store,
            artifact_bucket="artifacts",
        )
        matrix = await service.matrix(
            tenant_id=kwargs["tenant_id"], actor_id=kwargs["actor_id"], run_id=run.id
        )
        unit = matrix.units[0]
        revision = await service.create_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            expected_revision=1,
            idempotency_key="xlsx-value-v2",
            request=ManualRevisionInput(
                action="final",
                value='=HYPERLINK("https://example.invalid")',
                evidence_state="insufficient",
                reason="Confirmed directly by the buyer.",
                value_origin="user_provided",
            ),
        )
        with pytest.raises(AnalysisError, match="analysis_review_required"):
            await service.create_export(
                tenant_id=kwargs["tenant_id"],
                actor_id=kwargs["actor_id"],
                run_id=run.id,
                idempotency_key="reviewed-before-review",
                request=ExportInput(mode="reviewed"),
            )
        await service.review_cell_revision(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            unit_id=unit.id,
            idempotency_key="xlsx-approve-v2",
            request=ReviewInput(
                revision=2,
                content_hash=revision.content_hash,
                decision="approved",
                reason="Ready for the reviewed workbook.",
            ),
        )
        artifact = await service.create_export(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            run_id=run.id,
            idempotency_key="reviewed-xlsx-v1",
            request=ExportInput(mode="reviewed"),
        )
        replay = await service.create_export(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            run_id=run.id,
            idempotency_key="reviewed-xlsx-v1",
            request=ExportInput(mode="reviewed"),
        )
        assert replay.id == artifact.id
        assert artifact.status == "ready"
        assert artifact.cell_revision_ids == (revision.id,)
        second_artifact = await service.create_export(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            run_id=run.id,
            idempotency_key="reviewed-xlsx-v2",
            request=ExportInput(mode="reviewed"),
        )
        assert second_artifact.id != artifact.id
        assert second_artifact.matrix_revision == artifact.matrix_revision + 1
        assert second_artifact.matrix_hash == artifact.matrix_hash
        assert second_artifact.object_key != artifact.object_key
        assert (artifact.object_bucket, artifact.object_key) in store.objects
        assert (second_artifact.object_bucket, second_artifact.object_key) in store.objects
        body = store.objects[(artifact.object_bucket, artifact.object_key)][0]
        assert hashlib.sha256(body).hexdigest() == artifact.content_hash
        workbook = load_workbook(BytesIO(body), data_only=False, read_only=True)
        assert workbook.sheetnames == ["Comparison", "Pending Items", "Sources & Basis"]
        cells = [cell for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row]
        assert all(cell.data_type != "f" for cell in cells)
        dangerous = [cell.value for cell in cells if "HYPERLINK" in str(cell.value)]
        assert dangerous == ['\'=HYPERLINK("https://example.invalid")']
        assert workbook["Pending Items"].max_row == 2
        workbook.close()

        approval_input = ArtifactApprovalInput(
            artifact_hash=artifact.content_hash,
            source_manifest_hash=artifact.source_manifest_hash,
            matrix_hash=artifact.matrix_hash,
            operation="deliver",
            decision="approved",
            reason="Approved for delivery.",
            expires_at=datetime.now(UTC) + timedelta(minutes=30),
        )
        with pytest.raises(AnalysisError, match="analysis_approval_target_conflict"):
            await service.approve_artifact(
                tenant_id=kwargs["tenant_id"],
                actor_id=kwargs["actor_id"],
                artifact_id=artifact.id,
                idempotency_key="wrong-artifact-hash",
                request=approval_input.model_copy(update={"artifact_hash": "0" * 64}),
            )
        approval = await service.approve_artifact(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            artifact_id=artifact.id,
            idempotency_key="approve-artifact-v1",
            request=approval_input,
        )
        approval_replay = await service.approve_artifact(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            artifact_id=artifact.id,
            idempotency_key="approve-artifact-v1",
            request=approval_input,
        )
        assert approval_replay.id == approval.id
        assert approval.consumed_at is not None
        download = await service.download_artifact(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            artifact_id=artifact.id,
        )
        assert download.url.endswith(artifact.object_key)
        assert download.content_hash == artifact.content_hash
        assert store.presigns == 1

        async with browser_db.sessions.begin() as session:
            document = await session.get(Document, document_id)
            assert document is not None
            document.access_mode = "restricted"
        with pytest.raises(AnalysisError, match="analysis_not_found"):
            await service.download_artifact(
                tenant_id=kwargs["tenant_id"],
                actor_id=kwargs["actor_id"],
                artifact_id=artifact.id,
            )
        assert store.presigns == 1


async def test_export_replay_recovers_after_object_write_and_database_rollback(browser_db) -> None:
    store = MemoryArtifactStore()
    async with httpx.AsyncClient(transport=httpx.MockTransport(response)) as provider:
        execution, kwargs, _ = await setup(browser_db.sessions, provider)
        run = await execution.create_run(**kwargs)
        runtime, claim = await claim_run(browser_db.sessions, run)
        await execution.execute(claim)
        await runtime.succeed(claim)

        interrupted = AnalysisReviewDeliveryService(
            session_factory=RollbackSecondTransaction(browser_db.sessions),
            gateway=execution.gateway,
            artifact_store=store,
        )
        with pytest.raises(RuntimeError, match="simulated finalization rollback"):
            await interrupted.create_export(
                tenant_id=kwargs["tenant_id"],
                actor_id=kwargs["actor_id"],
                run_id=run.id,
                idempotency_key="recover-written-object",
                request=ExportInput(mode="draft"),
            )

        async with browser_db.sessions() as session:
            stored = await session.scalar(select(AnalysisArtifact))
            assert stored is not None and stored.status == "writing"
            artifact_id = stored.id
            object_key = stored.object_key

        recovered = await AnalysisReviewDeliveryService(
            session_factory=browser_db.sessions,
            gateway=execution.gateway,
            artifact_store=store,
        ).create_export(
            tenant_id=kwargs["tenant_id"],
            actor_id=kwargs["actor_id"],
            run_id=run.id,
            idempotency_key="recover-written-object",
            request=ExportInput(mode="draft"),
        )
        assert recovered.id == artifact_id
        assert recovered.object_key == object_key
        assert recovered.status == "ready"
        async with browser_db.sessions() as session:
            assert await session.scalar(select(func.count()).select_from(AnalysisArtifact)) == 1
