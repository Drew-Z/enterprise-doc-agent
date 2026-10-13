from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from enterprise_doc_core.documents.ingestion import DocumentParseViolation
from enterprise_doc_core.documents.structure import StructureLimits, parse_document_structure


def test_docx_preserves_body_order_merged_cells_and_source_positions() -> None:
    document = Document()
    document.add_heading("报价", level=1)
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "实施费"
    table.cell(1, 1).text = "1200 元"
    table.cell(2, 0).merge(table.cell(2, 1)).text = "税费另计"
    document.add_paragraph("交付时间为收到材料后十个工作日。")
    buffer = BytesIO()
    document.save(buffer)

    parsed = parse_document_structure(buffer.getvalue(), extension=".docx")
    texts = [block.text for block in parsed.blocks]
    assert texts == [
        "报价",
        "项目",
        "金额",
        "实施费",
        "1200 元",
        "税费另计",
        "交付时间为收到材料后十个工作日。",
    ]
    price = next(block for block in parsed.blocks if block.text == "1200 元")
    assert price.locator.kind == "docx"
    assert price.locator.row == 2 and price.locator.column == 2
    assert price.locator.page is None
    merged = next(block for block in parsed.blocks if block.text == "税费另计")
    assert merged.locator.column_span == 2
    assert parsed.coverage == "complete"
    assert parse_document_structure(buffer.getvalue(), extension="docx") == parsed
    for block in parsed.blocks:
        assert parsed.text[block.start_offset : block.end_offset] == block.text


def test_xlsx_keeps_hidden_sources_and_flags_uncomputed_formula() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet.append(["数量", "单价", "总价"])
    sheet.append([3, 120, "=A2*B2"])
    sheet["B2"].number_format = '"¥"0.00'
    extra = workbook.create_sheet("补充条款")
    extra.sheet_state = "hidden"
    extra["A1"] = "未包含安装服务"
    buffer = BytesIO()
    workbook.save(buffer)

    parsed = parse_document_structure(buffer.getvalue(), extension=".xlsx")
    formula = next(block for block in parsed.blocks if block.locator.cell == "C2")
    assert formula.formula == "=A2*B2" and formula.cached_value is None
    assert any(issue.code == "formula_cache_missing" for issue in parsed.issues)
    assert parsed.coverage == "partial"
    hidden = next(block for block in parsed.blocks if block.text == "未包含安装服务")
    assert hidden.locator.sheet == "补充条款" and hidden.locator.hidden
    price = next(block for block in parsed.blocks if block.locator.cell == "B2")
    assert price.number_format == '"¥"0.00'
    assert not any(block.text == "360" for block in parsed.blocks)


def test_blank_pdf_reports_ocr_gap_instead_of_claiming_full_coverage() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    parsed = parse_document_structure(buffer.getvalue(), extension=".pdf")
    assert not parsed.blocks
    assert parsed.coverage == "partial"
    assert any(issue.code == "needs_ocr" and issue.locator.page == 1 for issue in parsed.issues)


def test_text_has_stable_line_locators_and_rejects_invalid_encoding() -> None:
    parsed = parse_document_structure(
        "# 范围\r\n支持五名用户\r\n税费未说明".encode(), extension="txt"
    )
    assert parsed.blocks[1].locator.line == 2
    assert parsed.blocks[1].heading == "范围"
    with pytest.raises(DocumentParseViolation, match="UTF-8"):
        parse_document_structure(b"\xff\x00", extension="txt")


def test_size_and_table_limits_fail_explicitly() -> None:
    with pytest.raises(DocumentParseViolation) as error:
        parse_document_structure(b"four", extension="txt", limits=StructureLimits(max_file_bytes=3))
    assert error.value.code == "structure_file_limit"
    workbook = Workbook()
    workbook.active.append([1, 2, 3])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(DocumentParseViolation) as error:
        parse_document_structure(
            buffer.getvalue(), extension="xlsx", limits=StructureLimits(max_cells=2)
        )
    assert error.value.code == "structure_cell_limit"


def test_office_archive_with_unsafe_member_is_rejected_before_parsing() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("xl/workbook.xml", "<workbook />")
        archive.writestr("../escape.xml", "untrusted")
    with pytest.raises(DocumentParseViolation) as error:
        parse_document_structure(buffer.getvalue(), extension="xlsx")
    assert error.value.code == "document_xlsx_path_unsafe"


def test_xlsx_external_link_is_reported_as_unsupported() -> None:
    workbook = Workbook()
    workbook.active["A1"] = "公开价格"
    source = BytesIO()
    workbook.save(source)
    workbook.close()

    rebuilt = BytesIO()
    with (
        ZipFile(BytesIO(source.getvalue())) as original,
        ZipFile(rebuilt, "w", ZIP_DEFLATED) as archive,
    ):
        for entry in original.infolist():
            archive.writestr(entry, original.read(entry.filename))
        archive.writestr(
            "xl/externalLinks/externalLink1.xml",
            '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
        )

    parsed = parse_document_structure(rebuilt.getvalue(), extension="xlsx")
    assert parsed.coverage == "partial"
    assert any(issue.code == "xlsx_unsupported_content" for issue in parsed.issues)


def test_corrupt_office_archive_fails_with_typed_error() -> None:
    with pytest.raises(DocumentParseViolation) as error:
        parse_document_structure(b"not-an-office-archive", extension="xlsx")
    assert error.value.code == "document_xlsx_zip_invalid"


def test_office_archive_entry_limit_fails_before_parsing() -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types />")
        archive.writestr("xl/workbook.xml", "<workbook />")
        for index in range(1999):
            archive.writestr(f"xl/dummy/{index}.xml", "<item />")
    with pytest.raises(DocumentParseViolation) as error:
        parse_document_structure(buffer.getvalue(), extension="xlsx")
    assert error.value.code == "document_xlsx_entry_limit_exceeded"


async def test_isolated_parser_returns_the_same_canonical_structure() -> None:
    from enterprise_doc_core.documents.structure_process import parse_document_isolated

    content = "支持五名用户".encode()
    assert await parse_document_isolated(content, extension="txt") == parse_document_structure(
        content, extension="txt"
    )


async def test_isolated_parser_timeout_does_not_prevent_a_later_parse() -> None:
    from enterprise_doc_core.documents.structure_process import parse_document_isolated

    with pytest.raises(DocumentParseViolation) as error:
        await parse_document_isolated(b"source", extension="txt", timeout_seconds=0.001)
    assert error.value.code == "structure_timeout"
    assert (await parse_document_isolated(b"source", extension="txt")).blocks[0].text == "source"


def test_nested_docx_table_keeps_paragraphs_on_both_sides_in_order() -> None:
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "before"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "nested"
    cell.add_paragraph("after")
    buffer = BytesIO()
    document.save(buffer)
    parsed = parse_document_structure(buffer.getvalue(), extension="docx")
    assert [block.text for block in parsed.blocks] == ["before", "nested", "after"]
    assert parsed.blocks[0].locator.table == parsed.blocks[2].locator.table
    assert parsed.blocks[0].locator.paragraph != parsed.blocks[2].locator.paragraph


def test_xlsx_grouped_hidden_columns_are_all_marked_hidden() -> None:
    workbook = Workbook()
    workbook.active.append(["public", "hidden B", "hidden C", "hidden D"])
    workbook.active.column_dimensions.group("B", "D", hidden=True)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    parsed = parse_document_structure(buffer.getvalue(), extension="xlsx")
    assert [block.locator.hidden for block in parsed.blocks] == [False, True, True, True]


async def test_cancelling_an_active_parser_reaps_its_process(monkeypatch) -> None:
    import asyncio
    import subprocess

    from enterprise_doc_core.documents.structure_process import parse_document_isolated

    started = asyncio.Event()
    loop = asyncio.get_running_loop()
    processes = []
    original = subprocess.Popen

    def observe(*args, **kwargs):
        process = original(*args, **kwargs)
        processes.append(process)
        loop.call_soon_threadsafe(started.set)
        return process

    monkeypatch.setattr(subprocess, "Popen", observe)
    task = asyncio.create_task(parse_document_isolated(b"source", extension="txt"))
    await started.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert processes and processes[0].poll() is not None
