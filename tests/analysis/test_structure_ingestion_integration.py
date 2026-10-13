from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from sqlalchemy import func, select

from enterprise_doc_core.documents.ingestion_service import (
    DocumentIngestionError,
    DocumentIngestionService,
    IngestionVersions,
)
from enterprise_doc_core.documents.models import DocumentChunk, DocumentVersion
from enterprise_doc_core.documents.structure import ParsedDocumentStructure
from enterprise_doc_core.documents.structure_models import DocumentStructure
from enterprise_doc_core.jobs import JobRuntimeService, RetryDisposition, create_job_records
from tests.jobs.test_m3_ingestion_integration import (
    FailOnceEmbeddingProvider,
    FakeObjectStore,
    _seed_uploaded_document,
)

pytestmark = pytest.mark.integration


async def test_structured_ingestion_recovers_persisted_quotes_without_reparsing(browser_db) -> None:
    document = Document()
    document.add_paragraph("费用说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "实施费"
    table.cell(0, 1).text = "1200 元"
    buffer = BytesIO()
    document.save(buffer)
    content = buffer.getvalue()
    tenant, actor, version, _ = await _seed_uploaded_document(browser_db.sessions, content=content)
    async with browser_db.sessions.begin() as session:
        row = await session.get(DocumentVersion, version)
        row.original_filename = "quote.docx"
        created = await create_job_records(
            session,
            tenant_id=tenant,
            actor_id=actor,
            job_type="document.ingest",
            idempotency_key="structured-ingestion",
            document_version_id=version,
            payload={"document_version_id": str(version)},
        )
    runtime = JobRuntimeService(session_factory=browser_db.sessions, retry_base_seconds=0.001)
    claim = await runtime.claim(job_id=created.job_id, worker_id="analysis-test")
    assert claim is not None
    objects = FakeObjectStore(content)
    embedding = FailOnceEmbeddingProvider()
    service = DocumentIngestionService(
        session_factory=browser_db.sessions,
        object_store=objects,
        documents_bucket="documents",
        embedding_provider=embedding,
        versions=IngestionVersions(parser=2),
    )
    with pytest.raises(DocumentIngestionError):
        await service(claim)
    async with browser_db.sessions() as session:
        snapshot = await session.scalar(select(DocumentStructure))
        assert snapshot is not None
        parsed = ParsedDocumentStructure.model_validate(snapshot.payload)
        assert [block.text for block in parsed.blocks] == ["费用说明", "实施费", "1200 元"]
        original_id = snapshot.generation_id
        original_hash = snapshot.payload_sha256
    await runtime.fail(
        claim,
        error_code="test_embedding_failure",
        error_message="Controlled failure",
        disposition=RetryDisposition.RETRYABLE,
    )
    import asyncio

    await asyncio.sleep(0.01)
    resumed = await runtime.claim(job_id=created.job_id, worker_id="analysis-recovery")
    assert resumed is not None
    await service(resumed)
    await runtime.succeed(resumed)
    assert objects.head_calls == 1
    async with browser_db.sessions() as session:
        snapshot = await session.get(DocumentStructure, original_id)
        assert snapshot.payload_sha256 == original_hash
        assert await session.scalar(select(func.count()).select_from(DocumentStructure)) == 1
        chunks = (await session.scalars(select(DocumentChunk))).all()
        assert any(row.normalized_text == "1200 元" for row in chunks)
        for row in chunks:
            assert parsed.text[row.start_offset : row.end_offset] == row.normalized_text
